"""Générateurs de documents SYNTHÉTIQUES pour les tests (aucun vrai entretien).

Le PDF est construit à la main (police Helvetica standard, WinAnsiEncoding)
pour éviter une dépendance de test supplémentaire.
"""

import io

SAMPLE_LINES = [
    "Entretien synthétique n°1 — test",
    "",
    "Enquêteur : Bonjour, est-ce que tu peux te présenter ?",
    "Enquêté : Euh… oui. Alors euh je je suis en L2 de socio.",
    "J’utilise ChatGPT, genre : pour tout, enfin presque.",
    "Enquêteur : Et à 14:30, pendant le TD, tu l'utilises ?",
    "Enquêté : Bah ouais, ouais, c'est chiant mais bon, j'avoue.",
]


def make_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row, values in zip(table.rows, table_rows):
            for cell, value in zip(row.cells, values):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _escape(text: str) -> bytes:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").encode("cp1252")


def make_pdf(pages: list[list[str]]) -> bytes:
    """PDF textuel minimal : une liste de lignes par page (liste vide = page blanche)."""
    objects: list[bytes | None] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    objects.append(None)  # arbre des pages, rempli plus bas
    pages_id = len(objects)
    kids = []
    for lines in pages:
        stream = b""
        if lines:
            stream = b"BT /F1 11 Tf 14 TL 50 800 Td " + b"".join(b"(" + _escape(l) + b") Tj T* " for l in lines) + b"ET"
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        kids.append(add(
            b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>" % (pages_id, font, content)
        ))
    objects[pages_id - 1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
        b" ".join(b"%d 0 R" % k for k in kids), len(kids))
    catalog = add(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id)

    out, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, catalog, xref)
    return out
